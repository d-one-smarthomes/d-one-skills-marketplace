#!/usr/bin/env python3
"""
Estimator - build a detailed D-One quote (equipment + labour) from a list of items.

Usage:
    python build_quote.py items.json output.xlsx [odoo_import.csv]

If the third argument is omitted, the Odoo import file is written next to the
xlsx with the same stem and an "_OdooImport.csv" suffix.

items.json schema:
{
  "project_name": "Smith Residence",
  "quote_ref": "Q-1234",           # optional -> Order Reference in Odoo
  "customer": "Smith Family Trust", # optional -> Customer (partner) in Odoo import
  "items": [
    {"query": "IPMX-E20F-IRB2", "qty": 4, "area": "Ground Floor", "room": "Lounge", "system": "Multiroom Audio"},
    {"query": "ICREALTIME 2 Megapixel IP Camera", "qty": 2, "room": "Front Door"}
  ]
}

"area", "room" and "system" are optional per item. Each has its own Excel column and
its own Odoo order-line field: Area -> x_studio_area, Room -> x_room,
System -> x_system. System must be one of Odoo's System types (label or key); if
omitted it is pre-filled from the System that product was most often quoted under
in Odoo. Labour lines never carry an Area, Room or System.

"query" can be a SKU (preferred, exact match) or a free-text product name/description.

PRICING SOURCE (this is the important bit):
  Equipment SELL price is pulled LIVE from Odoo (product.product -> list_price)
  whenever Odoo is reachable, so quotes always reflect the current price in
  D-One's business system. Odoo is reached read-only via scripts/odoo_client.py,
  using ODOO_URL / ODOO_DB / ODOO_LOGIN / ODOO_KEY from the environment.

  If Odoo is NOT reachable (no key in the environment, offline, auth fails), the
  script transparently falls back to the local inventory.csv snapshot and clearly
  flags every line whose price came from the snapshot rather than live Odoo.

  Cost history (equip_cost_by_sku.json), labour hours (labour_hours_by_*.json) and
  the rate card (rate_card.json) are always local — Odoo holds the sell price, not
  D-One's install-hour history or per-SKU cost.

OUTPUTS:
  1. <output>.xlsx        — the human-facing D-One quote (equipment + labour).
  2. <output>_OdooImport.csv — the SAME quote in Odoo's Sales Order import format,
     ready to import straight into Odoo once the quote is confirmed. Only lines
     that resolve to a real Odoo product id are written to this file, so the import
     is clean; anything that couldn't be mapped is listed in the run summary and
     stays in the xlsx for a human to handle.
"""
import sys
import os
import csv
import json
import difflib

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# odoo_client lives next to this script
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
try:
    import odoo_client
except Exception:  # pragma: no cover - defensive
    odoo_client = None

ASSETS = os.path.join(os.path.dirname(__file__), "..", "assets")

# D-One brand palette - matches conduit-schedule / component-schedule skills
NAVY = "1B2A41"
DONE_BLUE = "1379C9"
MID_BLUE = "5FA8D3"
LIGHT_ROW = "EAF3FB"
WHITE = "FFFFFF"
FLAG_AMBER = "FDECC8"
FLAG_RED = "F8D7DA"


# --------------------------------------------------------------------------- #
# Local reference data
# --------------------------------------------------------------------------- #
def load_reference():
    inventory = {}
    inventory_by_name = []
    with open(os.path.join(ASSETS, "inventory.csv"), newline="", encoding="utf-8") as f:
        for row in csv.DictReader(f):
            sku = (row["sku"] or "").strip()
            item = {
                "sku": sku,
                "name": row["name"],
                "category": row["category"] or "Uncategorised",
                "price": float(row["price_zar"] or 0),
                "unit": row["unit"] or "Units",
            }
            if sku:
                inventory[sku.upper()] = item
            inventory_by_name.append(item)

    with open(os.path.join(ASSETS, "equip_cost_by_sku.json")) as f:
        equip_cost = json.load(f)
    with open(os.path.join(ASSETS, "labour_hours_by_sku.json")) as f:
        labour_sku = json.load(f)
    with open(os.path.join(ASSETS, "labour_hours_by_category.json")) as f:
        labour_cat = json.load(f)
    with open(os.path.join(ASSETS, "rate_card.json")) as f:
        rates = json.load(f)

    return inventory, inventory_by_name, equip_cost, labour_sku, labour_cat, rates


def match_csv_item(query, inventory, inventory_by_name):
    """CSV fallback match. Return (item, match_type) where match_type is 'sku','name',None."""
    q = query.strip()
    if q.upper() in inventory:
        return inventory[q.upper()], "sku"
    names = [i["name"] for i in inventory_by_name]
    best = difflib.get_close_matches(q, names, n=1, cutoff=0.6)
    if best:
        for i in inventory_by_name:
            if i["name"] == best[0]:
                return i, "name"
    return None, None


def csv_category_for_sku(sku, inventory):
    if sku and sku.upper() in inventory:
        return inventory[sku.upper()]["category"]
    return "Uncategorised"


def get_labour_hours(sku, category, labour_sku, labour_cat):
    if sku:
        key = sku.strip().upper()
        for k, v in labour_sku.items():
            if k.strip().upper() == key:
                return v["fix1_hours"], v["fix2_hours"], v["programming_hours"], "history (this exact item)"
    if category in labour_cat:
        v = labour_cat[category]
        return v["fix1_hours"], v["fix2_hours"], v["programming_hours"], f"category average ({category})"
    return 0, 0, 0, "no data - needs manual hours"


def cost_for(sku, unit_price, equip_cost, rates):
    if sku:
        key = sku.strip().upper()
        for k, v in equip_cost.items():
            if k.strip().upper() == key:
                return v["unit_cost"], "history"
    if unit_price:
        return round(unit_price / rates["equipment_fallback_markup"], 2), "estimated (no cost history for this SKU)"
    return 0, "no data"


# --------------------------------------------------------------------------- #
# Odoo live pricing
# --------------------------------------------------------------------------- #
def fetch_odoo_prices(queries):
    """
    Try to resolve every query against live Odoo. Returns (odoo_map, status) where
    odoo_map maps query -> {"id","sku","name","price","match_type"} (only for hits)
    and status is a short human string describing what happened.
    """
    if odoo_client is None or not odoo_client.available():
        return {}, "Odoo skipped — no ODOO_KEY in environment; using inventory.csv snapshot."
    try:
        od = odoo_client.Odoo()
    except Exception as e:  # OdooError or anything unexpected
        return {}, f"Odoo unreachable ({e}) — using inventory.csv snapshot."

    odoo_map = {}
    for q in queries:
        if q in odoo_map:
            continue
        try:
            prod, mtype = odoo_client.lookup_product(od, q)
        except Exception:
            prod, mtype = None, None
        if prod:
            odoo_map[q] = {
                "id": prod["id"],
                "sku": (prod.get("default_code") or "").strip(),
                "name": prod.get("name") or q,
                "price": float(prod.get("list_price") or 0),
                "match_type": mtype,
            }
    return odoo_map, f"Odoo live: priced {len(odoo_map)}/{len(set(queries))} unique items from list_price."


def fetch_system_info(product_ids):
    """
    Live from Odoo: the System choices (sale.order.line.x_system) and, per product,
    how often past quote lines used each System. Returns (options, history, note)
    where options = [(key, label), ...]. Products themselves carry no System in
    Odoo, so history is how a new line's System gets pre-filled.
    """
    if odoo_client is None or not odoo_client.available():
        return [], {}, "System list unavailable (no Odoo) — System column left blank."
    try:
        od = odoo_client.Odoo()
        options = odoo_client.system_options(od)
        history = odoo_client.system_history(od, product_ids)
    except Exception as e:
        return [], {}, f"System list unavailable ({e}) — System column left blank."
    return options, history, (f"System: {len(options)} types read live from Odoo; "
                              f"history found for {len(history)} product(s).")


def normalise_system(value, options):
    """Map a System given as label or key (any case) to (key, label). None if not valid."""
    v = (value or "").strip().lower()
    if not v:
        return None
    for k, lbl in options:
        if v == str(k).lower() or v == str(lbl).lower():
            return k, lbl
    return None


def assign_system(line, explicit, options, history):
    """Set line['system_key'] / line['system'] (label) and add a flag if it needs a look.
    Priority: explicit value from items.json -> most-used System in Odoo history -> blank."""
    labels = dict(options)
    line["system_key"], line["system"] = "", ""
    note = None
    if explicit:
        hit = normalise_system(explicit, options)
        if hit:
            line["system_key"], line["system"] = hit
            return
        note = f'System "{explicit}" is not one of the Odoo System types — choose from the dropdown'
    elif not options:
        return
    else:
        counts = history.get(line.get("odoo_id")) or {}
        counts = {k: n for k, n in counts.items() if k in labels}
        if counts:
            ranked = sorted(counts.items(), key=lambda kv: -kv[1])
            key = ranked[0][0]
            line["system_key"], line["system"] = key, labels[key]
            if len(ranked) > 1:
                total = sum(counts.values())
                others = ", ".join(f"{labels[k]} ({n})" for k, n in ranked[1:])
                note = (f"System pre-filled as {labels[key]} ({ranked[0][1]} of {total} past lines) "
                        f"— also quoted as {others}; confirm")
        else:
            note = "No System history for this product in Odoo — choose System from the dropdown"
    if note:
        line["flag"] = " | ".join(x for x in (line.get("flag"), note) if x)


def resolve_labour_products(rates):
    """
    Resolve the three labour phases to Odoo product ids using rate_card ->
    odoo_labour_products. Returns (map phase->{"id","name"}, note). Empty if Odoo
    unavailable or none configured/resolvable.
    """
    cfg = (rates.get("odoo_labour_products") or {})
    wanted = {k: cfg.get(k, "") for k in ("fix1_cabling", "fix2_installation", "programming")}
    if not any(wanted.values()):
        return {}, "labour products not configured (odoo_labour_products blank) — labour left out of import."
    if odoo_client is None or not odoo_client.available():
        return {}, "labour products not resolved — Odoo unavailable."
    try:
        od = odoo_client.Odoo()
    except Exception:
        return {}, "labour products not resolved — Odoo unreachable."
    out = {}
    unresolved = []
    for phase, ref in wanted.items():
        if not ref:
            continue
        try:
            prod = odoo_client.resolve_service_product(od, ref)
        except Exception:
            prod = None
        if prod:
            out[phase] = {"id": prod["id"], "name": prod.get("name") or ref}
        else:
            unresolved.append(f"{phase}='{ref}'")
    note = f"labour products resolved: {len(out)}/{sum(1 for v in wanted.values() if v)}."
    if unresolved:
        note += " unresolved: " + ", ".join(unresolved)
    return out, note


# --------------------------------------------------------------------------- #
# Line computation
# --------------------------------------------------------------------------- #
def compute_line(query, qty, odoo_hit, inventory, inventory_by_name, equip_cost, labour_sku, labour_cat, rates):
    sku = name = None
    unit_price = 0.0
    odoo_id = None
    price_source = None
    match_type = None
    flag_bits = []

    if odoo_hit:
        odoo_id = odoo_hit["id"]
        sku = odoo_hit["sku"] or ""
        name = odoo_hit["name"]
        unit_price = odoo_hit["price"]
        match_type = odoo_hit["match_type"]
        price_source = "Odoo live (list_price)"
        if match_type == "name":
            flag_bits.append(f'Odoo matched by NAME, not SKU — verify this is the right product: "{name}"')
    else:
        item, match_type = match_csv_item(query, inventory, inventory_by_name)
        if not item:
            return {
                "query": query, "matched_name": None, "sku": None, "category": "UNMATCHED",
                "qty": qty, "unit_price": 0, "unit_cost": 0, "cost_source": None,
                "fix1_h": 0, "fix2_h": 0, "prog_h": 0, "hours_source": None,
                "odoo_id": None, "price_source": None,
                "flag": "NOT FOUND in Odoo or inventory.csv — price and hours must be entered manually",
            }
        sku = item["sku"]
        name = item["name"]
        unit_price = item["price"]
        price_source = "inventory.csv snapshot (not live Odoo)"
        flag_bits.append("Priced from local snapshot, not live Odoo — confirm the current Odoo price")
        if match_type == "name":
            flag_bits.append(f'Matched by name, not SKU — verify this is the right product: "{name}"')

    category = csv_category_for_sku(sku, inventory)
    unit_cost, cost_source = cost_for(sku, unit_price, equip_cost, rates)
    fix1_h, fix2_h, prog_h, hours_source = get_labour_hours(sku, category, labour_sku, labour_cat)
    if hours_source == "no data - needs manual hours":
        flag_bits.append("No labour history for this item or category — hours need manual entry")

    return {
        "query": query, "matched_name": name, "sku": sku, "category": category,
        "qty": qty, "unit_price": unit_price, "unit_cost": unit_cost, "cost_source": cost_source,
        "fix1_h": fix1_h, "fix2_h": fix2_h, "prog_h": prog_h, "hours_source": hours_source,
        "odoo_id": odoo_id, "price_source": price_source,
        "flag": " | ".join(flag_bits) if flag_bits else None,
    }


# --------------------------------------------------------------------------- #
# Excel output
# --------------------------------------------------------------------------- #
def build_workbook(lines, project_name, quote_ref, rates, pricing_note, system_labels=()):
    wb = Workbook()
    ws = wb.active
    ws.title = "Quote"

    r = rates["labour_rates_per_hour"]
    fix1_rate, fix2_rate, prog_rate = r["fix1_cabling"]["sell"], r["fix2_installation"]["sell"], r["programming"]["sell"]

    # Columns. Area, Room and System come first so the quote reads area-by-area,
    # room-by-room (how installs are planned). Each maps to its own Odoo order-line
    # field: Area -> x_studio_area, Room -> x_room, System -> x_system.
    # Money and totals columns are indexed off this order — if you reorder, update
    # the COL_* constants below to match.
    headers = ["Area", "Room", "System", "Category", "Item", "SKU", "Qty", "Unit Price", "Unit Cost",
               "Equip Total", "Equip Cost Total", "1st Fix Hrs", "2nd Fix Hrs", "Prog Hrs",
               "Labour Total", "Price Source", "Notes"]
    COL_SYSTEM = 3
    COL_UNIT_PRICE, COL_UNIT_COST = 8, 9
    COL_EQUIP_TOTAL, COL_EQUIP_COST, COL_LABOUR_TOTAL = 10, 11, 15
    money_cols = (COL_UNIT_PRICE, COL_UNIT_COST, COL_EQUIP_TOTAL, COL_EQUIP_COST, COL_LABOUR_TOTAL)
    ncols = len(headers)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    title_cell = ws.cell(row=1, column=1, value=f"D-One Quote — {project_name}" + (f" ({quote_ref})" if quote_ref else ""))
    title_cell.font = Font(color=WHITE, size=14, bold=True)
    title_cell.fill = PatternFill("solid", fgColor=NAVY)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26

    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=2, column=c, value=h)
        cell.font = Font(color=WHITE, bold=True)
        cell.fill = PatternFill("solid", fgColor=DONE_BLUE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    row = 3
    grand_equip_price = grand_equip_cost = grand_labour_price = 0
    # sort area, then room, then category; blanks sort last
    first_data_row = row
    for line in sorted(lines, key=lambda l: ((l.get("area") or "~~~").lower(),
                                             (l.get("room") or "~~~").lower(), l["category"])):
        equip_total = line["unit_price"] * line["qty"]
        equip_cost_total = line["unit_cost"] * line["qty"]
        labour_total = (line["fix1_h"] * fix1_rate + line["fix2_h"] * fix2_rate + line["prog_h"] * prog_rate) * line["qty"]

        grand_equip_price += equip_total
        grand_equip_cost += equip_cost_total
        grand_labour_price += labour_total

        values = [
            line.get("area", ""), line.get("room", ""), line.get("system", ""),
            line["category"], line["matched_name"] or line["query"],
            line["sku"] or "-", line["qty"], line["unit_price"], line["unit_cost"],
            equip_total, equip_cost_total,
            line["fix1_h"], line["fix2_h"], line["prog_h"], round(labour_total, 2),
            line["price_source"] or "-", line["flag"] or "",
        ]
        fill = None
        if line["category"] == "UNMATCHED":
            fill = FLAG_RED
        elif line["flag"]:
            fill = FLAG_AMBER
        elif row % 2 == 0:
            fill = LIGHT_ROW
        for c, v in enumerate(values, start=1):
            cell = ws.cell(row=row, column=c, value=v)
            if fill:
                cell.fill = PatternFill("solid", fgColor=fill)
            if c in money_cols:
                cell.number_format = '#,##0.00'
        row += 1

    last_data_row = row - 1

    # System dropdown: limited to Odoo's own System types (read live), held on a
    # hidden "Lists" tab. Extends a little past the last line so added rows get it too.
    if system_labels:
        lists = wb.create_sheet("Lists")
        for i, lbl in enumerate(system_labels, start=1):
            lists.cell(row=i, column=1, value=lbl)
        lists.sheet_state = "hidden"
        dv = DataValidation(type="list", formula1=f"=Lists!$A$1:$A${len(system_labels)}",
                            allow_blank=True, showErrorMessage=True,
                            errorTitle="Not an Odoo System",
                            error="Pick a System from the list — it must match Odoo's System types.")
        ws.add_data_validation(dv)
        col_letter = get_column_letter(COL_SYSTEM)
        dv.add(f"{col_letter}{first_data_row}:{col_letter}{max(last_data_row, first_data_row) + 50}")

    row += 1
    ws.cell(row=row, column=1, value="TOTALS").font = Font(bold=True)
    ws.cell(row=row, column=COL_EQUIP_TOTAL, value=round(grand_equip_price, 2)).font = Font(bold=True)
    ws.cell(row=row, column=COL_EQUIP_COST, value=round(grand_equip_cost, 2)).font = Font(bold=True)
    ws.cell(row=row, column=COL_LABOUR_TOTAL, value=round(grand_labour_price, 2)).font = Font(bold=True)
    for c in (COL_EQUIP_TOTAL, COL_EQUIP_COST, COL_LABOUR_TOTAL):
        ws.cell(row=row, column=c).number_format = '#,##0.00'
    row += 1
    margin = grand_equip_price - grand_equip_cost
    ws.cell(row=row, column=1, value="Equipment margin (price - cost)")
    ws.cell(row=row, column=COL_EQUIP_TOTAL, value=round(margin, 2)).number_format = '#,##0.00'
    row += 1
    ws.cell(row=row, column=1, value="GRAND TOTAL (equip + labour, ex VAT)").font = Font(bold=True, size=12)
    ws.cell(row=row, column=COL_EQUIP_TOTAL, value=round(grand_equip_price + grand_labour_price, 2)).font = Font(bold=True, size=12)
    ws.cell(row=row, column=COL_EQUIP_TOTAL).number_format = '#,##0.00'
    row += 2
    note_cell = ws.cell(row=row, column=1, value=f"Pricing: {pricing_note}")
    note_cell.font = Font(italic=True, size=9)

    widths = [16, 18, 24, 20, 40, 16, 6, 11, 11, 12, 13, 10, 10, 9, 12, 26, 40]
    for c, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A3"
    return wb


# --------------------------------------------------------------------------- #
# Odoo Sales Order import CSV
# --------------------------------------------------------------------------- #
# Odoo imports a one-to-many (order + lines) by putting the parent (order) fields
# on the FIRST row of the record only, with a blank External ID ("id") on the
# continuation rows so they attach to the same order. Products are referenced by
# database id via the "/.id" suffix, which is unambiguous — no name matching.
IMPORT_HEADERS = [
    "id",                             # External ID — first row of the order only
    "name",                           # Order Reference — first row only
    "partner_id",                     # Customer — first row only
    "order_line/product_id/.id",      # product.product database id
    "order_line/name",                # line description
    "order_line/product_uom_qty",     # quantity
    "order_line/price_unit",          # unit price (ex VAT)
    "order_line/x_studio_area",       # Area (order line)
    "order_line/x_room",              # Room (order line)
    "order_line/x_system",            # System (order line) — Odoo selection key
]


def build_odoo_import(lines, project_name, quote_ref, customer, rates, labour_products):
    """
    Return (rows, skipped) for the Odoo import CSV.
    - Equipment lines with a real Odoo product id are written with product_id/.id.
    - Labour is aggregated into up to 3 service lines (one per phase) using the
      total hours across the quote, but only if that phase's Odoo product resolved.
    - Lines that can't be mapped cleanly are returned in `skipped` (not written),
      so the import file stays clean and importable.
    """
    r = rates["labour_rates_per_hour"]
    rate_by_phase = {
        "fix1_cabling": r["fix1_cabling"]["sell"],
        "fix2_installation": r["fix2_installation"]["sell"],
        "programming": r["programming"]["sell"],
    }
    phase_hours = {"fix1_cabling": 0.0, "fix2_installation": 0.0, "programming": 0.0}
    phase_label = {
        "fix1_cabling": "1st Fix Cabling",
        "fix2_installation": "2nd Fix Installation",
        "programming": "Programming",
    }

    rows = []
    skipped = []
    ext_id = "__import__.estimator_" + (quote_ref or project_name or "quote").strip().replace(" ", "_")

    first = True
    for line in lines:
        # accumulate labour hours for later aggregation
        phase_hours["fix1_cabling"] += line["fix1_h"] * line["qty"]
        phase_hours["fix2_installation"] += line["fix2_h"] * line["qty"]
        phase_hours["programming"] += line["prog_h"] * line["qty"]

        if not line.get("odoo_id"):
            skipped.append((line["matched_name"] or line["query"],
                            line.get("price_source") or "no Odoo product id"))
            continue
        rows.append({
            "id": ext_id if first else "",
            "name": (quote_ref or "") if first else "",
            "partner_id": (customer or "") if first else "",
            "order_line/product_id/.id": line["odoo_id"],
            "order_line/name": line["matched_name"] or line["query"],
            "order_line/product_uom_qty": line["qty"],
            "order_line/price_unit": round(line["unit_price"], 2),
            "order_line/x_studio_area": line.get("area", ""),
            "order_line/x_room": line.get("room", ""),
            "order_line/x_system": line.get("system_key", ""),
        })
        first = False

    # labour service lines
    for phase in ("fix1_cabling", "fix2_installation", "programming"):
        hrs = round(phase_hours[phase], 2)
        if hrs <= 0:
            continue
        prod = labour_products.get(phase)
        if not prod:
            skipped.append((f"Labour — {phase_label[phase]} ({hrs} hrs)",
                            "no Odoo service product mapped (set odoo_labour_products in rate_card.json)"))
            continue
        rows.append({
            "id": ext_id if first else "",
            "name": (quote_ref or "") if first else "",
            "partner_id": (customer or "") if first else "",
            "order_line/product_id/.id": prod["id"],
            "order_line/name": f"{phase_label[phase]} (labour)",
            "order_line/product_uom_qty": hrs,
            "order_line/price_unit": rate_by_phase[phase],
            # labour is aggregated across the whole quote: no Area, no Room, and
            # deliberately NO System (not even Professional Installation).
            "order_line/x_studio_area": "",
            "order_line/x_room": "",
            "order_line/x_system": "",
        })
        first = False

    return rows, skipped


def write_odoo_import(path, rows):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=IMPORT_HEADERS)
        w.writeheader()
        for row in rows:
            w.writerow(row)


# --------------------------------------------------------------------------- #
# Meta sheet — makes the Excel the editable source of truth
# --------------------------------------------------------------------------- #
# The Odoo import needs Customer and Quote Ref, which aren't line items. We store
# them on a "Meta" tab so a user can edit them in the Excel and the refresh step
# (refresh_import.py) reads them back. Labelled rows (label in col A, value in
# col B) so parsing is order-independent.
META_LABELS = {"project": "Project", "quote_ref": "Quote Ref", "customer": "Customer"}


def add_meta_sheet(wb, project_name, quote_ref, customer):
    ws = wb.create_sheet("Meta")
    title = ws.cell(row=1, column=1, value="D-One Estimator — Quote Info")
    title.font = Font(bold=True, size=12, color=WHITE)
    title.fill = PatternFill("solid", fgColor=NAVY)
    ws.merge_cells("A1:B1")
    data = [
        (META_LABELS["project"], project_name),
        (META_LABELS["quote_ref"], quote_ref),
        (META_LABELS["customer"], customer),
    ]
    for i, (label, val) in enumerate(data, start=2):
        ws.cell(row=i, column=1, value=label).font = Font(bold=True)
        ws.cell(row=i, column=2, value=val)
    note = ("Edit line items on the 'Quote' tab and Customer / Quote Ref above. "
            "After editing, regenerate the Odoo import from THIS file with "
            "refresh_import.py so the upload matches your edits — don't re-use the "
            "originally generated import CSV.")
    ws.cell(row=6, column=1, value="How to use").font = Font(bold=True)
    nc = ws.cell(row=6, column=2, value=note)
    nc.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 80
    return ws


# --------------------------------------------------------------------------- #
# Main
# --------------------------------------------------------------------------- #
def main():
    if len(sys.argv) not in (3, 4):
        print("Usage: python build_quote.py items.json output.xlsx [odoo_import.csv]")
        sys.exit(1)
    items_path, out_path = sys.argv[1], sys.argv[2]
    if len(sys.argv) == 4:
        import_path = sys.argv[3]
    else:
        stem = out_path[:-5] if out_path.lower().endswith(".xlsx") else out_path
        import_path = stem + "_OdooImport.csv"

    with open(items_path) as f:
        req = json.load(f)

    inventory, inventory_by_name, equip_cost, labour_sku, labour_cat, rates = load_reference()

    queries = [it["query"] for it in req["items"]]
    odoo_map, pricing_note = fetch_odoo_prices(queries)
    labour_products, labour_note = resolve_labour_products(rates)

    lines = []
    for it in req["items"]:
        odoo_hit = odoo_map.get(it["query"])
        line = compute_line(it["query"], it["qty"], odoo_hit, inventory, inventory_by_name,
                            equip_cost, labour_sku, labour_cat, rates)
        line["area"] = (it.get("area") or "").strip()
        line["room"] = (it.get("room") or "").strip()
        line["_system_in"] = (it.get("system") or "").strip()
        lines.append(line)

    sys_options, sys_history, system_note = fetch_system_info([l.get("odoo_id") for l in lines])
    for line in lines:
        assign_system(line, line.pop("_system_in"), sys_options, sys_history)

    project_name = req.get("project_name", "Untitled Project")
    quote_ref = req.get("quote_ref", "")
    customer = req.get("customer", "")

    wb = build_workbook(lines, project_name, quote_ref, rates, pricing_note,
                        [lbl for _k, lbl in sys_options])
    add_meta_sheet(wb, project_name, quote_ref, customer)
    wb.save(out_path)

    rows, skipped = build_odoo_import(lines, project_name, quote_ref, customer, rates, labour_products)
    write_odoo_import(import_path, rows)

    n_unmatched = sum(1 for l in lines if l["category"] == "UNMATCHED")
    n_flagged = sum(1 for l in lines if l["flag"] and l["category"] != "UNMATCHED")
    n_live = sum(1 for l in lines if (l.get("price_source") or "").startswith("Odoo live"))

    print(f"Saved {out_path} — {len(lines)} lines, {n_live} priced live from Odoo, "
          f"{n_unmatched} unmatched, {n_flagged} flagged for review")
    print(f"  {pricing_note}")
    print(f"  labour: {labour_note}")
    print(f"  {system_note}")
    print(f"Saved {import_path} — {len(rows)} importable Odoo lines, {len(skipped)} not written")
    if not customer:
        print("  NOTE: no 'customer' in items.json — set the Customer column in Odoo before/at import.")
    for name, reason in skipped:
        print(f"  SKIPPED from import: {name} — {reason}")


if __name__ == "__main__":
    main()
