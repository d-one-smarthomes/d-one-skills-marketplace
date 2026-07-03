#!/usr/bin/env python3
"""
D-One Conduit Schedule Builder
Usage: python build_conduit.py <data.json> <output.xlsx>

data.json schema:
{
  "project_name": "Price Drive",
  "quote_ref": "0179 R10",
  "sections": [
    {
      "level": "PERIMETER",
      "areas": [
        {
          "name": "Front Gate",
          "items": [
            {
              "point": "Door Intercom (UA-G3-Intercom)",
              "conduit": "25",
              "destination": "Head-End",
              "cable": "CAT6",
              "backbox": "Flush round outdoor",
              "power": "No"
            }
          ]
        }
      ]
    }
  ]
}
"""
import sys, json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# ── Colours ────────────────────────────────────────────────────────────────────
C_DARK  = "1F3864"   # dark navy  — title + column headers
C_LEVEL = "1379C9"   # D-One blue — level (PERIMETER / L1 BASEMENT …)
C_AREA  = "2E75B6"   # mid blue   — room sub-header
C_WHITE = "FFFFFF"
C_LGREY = "F2F2F2"   # alternating row tint
C_BORD  = "BFBFBF"

def _fill(hex_): return PatternFill("solid", start_color=hex_, end_color=hex_)
def _border():
    s = Side(style="thin", color=C_BORD)
    return Border(left=s, right=s, top=s, bottom=s)

def build(data: dict, out_path: str):
    wb = Workbook()
    ws = wb.active
    ws.title = "Conduit Schedule"

    project = data.get("project_name", "Project")
    ref     = data.get("quote_ref", "")
    title_str = f"CONDUIT SCHEDULE  —  {project}   |   Quote Ref: {ref}"

    # Column widths
    for col, w in zip("ABCDEF", [44, 10, 24, 28, 22, 10]):
        ws.column_dimensions[col].width = w

    # Row 1 — title bar
    ws.row_dimensions[1].height = 18
    ws.merge_cells("A1:F1")
    c = ws["A1"]
    c.value = title_str
    c.font      = Font(name="Arial", bold=True, size=13, color=C_WHITE)
    c.fill      = _fill(C_DARK)
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border    = _border()

    # Row 2 — column headers
    ws.row_dimensions[2].height = 16
    for ci, h in enumerate(["Point", "Conduit", "Destination", "Cable", "Backbox", "Power"], 1):
        c = ws.cell(row=2, column=ci)
        c.value     = h
        c.font      = Font(name="Arial", bold=True, size=9, color=C_WHITE)
        c.fill      = _fill(C_DARK)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border    = _border()

    row_idx   = 3
    item_alt  = 0   # alternating counter for data rows

    for section in data.get("sections", []):
        # Level header
        ws.row_dimensions[row_idx].height = 17
        ws.merge_cells(f"A{row_idx}:F{row_idx}")
        c = ws.cell(row=row_idx, column=1)
        c.value     = section["level"]
        c.font      = Font(name="Arial", bold=True, size=10, color=C_WHITE)
        c.fill      = _fill(C_LEVEL)
        c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        c.border    = _border()
        for ci in range(2, 7):
            ws.cell(row=row_idx, column=ci).border = _border()
        row_idx += 1

        for area in section.get("areas", []):
            # Area sub-header
            ws.row_dimensions[row_idx].height = 15
            ws.merge_cells(f"A{row_idx}:F{row_idx}")
            c = ws.cell(row=row_idx, column=1)
            c.value     = area["name"]
            c.font      = Font(name="Arial", bold=True, size=9, color=C_WHITE)
            c.fill      = _fill(C_AREA)
            c.alignment = Alignment(horizontal="left", vertical="center", indent=2)
            c.border    = _border()
            for ci in range(2, 7):
                ws.cell(row=row_idx, column=ci).border = _border()
            row_idx += 1

            for item in area.get("items", []):
                bg = C_LGREY if item_alt % 2 == 0 else C_WHITE
                ws.row_dimensions[row_idx].height = 14
                vals = [
                    f"■  {item.get('point', '')}",
                    item.get("conduit", "25"),
                    item.get("destination", ""),
                    item.get("cable", ""),
                    item.get("backbox", ""),
                    item.get("power", ""),
                ]
                for ci, v in enumerate(vals, 1):
                    c = ws.cell(row=row_idx, column=ci)
                    c.value  = v
                    c.font   = Font(name="Arial", size=9)
                    c.fill   = _fill(bg)
                    c.border = _border()
                    if ci == 1:
                        c.alignment = Alignment(horizontal="left",   vertical="center", indent=1)
                    elif ci == 2:
                        c.alignment = Alignment(horizontal="center", vertical="center")
                    else:
                        c.alignment = Alignment(horizontal="left",   vertical="center", indent=1)
                row_idx  += 1
                item_alt += 1

    # Page setup — A4 landscape, repeat headers
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize   = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth  = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.fitToPage   = True
    ws.print_title_rows       = "1:2"
    ws.freeze_panes           = "A3"

    wb.save(out_path)
    print(f"Saved → {out_path}  ({row_idx - 3} rows)")

if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Usage: python build_conduit.py data.json output.xlsx")
        sys.exit(1)
    with open(sys.argv[1]) as f:
        data = json.load(f)
    build(data, sys.argv[2])
