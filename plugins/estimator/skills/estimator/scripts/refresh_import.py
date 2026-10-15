#!/usr/bin/env python3
"""
Regenerate the Odoo Sales Order import file FROM an edited Excel quote.

Usage:
    python refresh_import.py <edited_quote.xlsx> [odoo_import.csv]

Why this exists:
    build_quote.py produces two files at once — the human-facing Excel quote and
    an Odoo import CSV. Once someone edits the Excel (changes a quantity or price,
    adds or removes a line, tweaks labour hours), that original CSV is stale.
    Run THIS script on the edited Excel right before uploading, and it rebuilds the
    import file so the upload matches the current state of the Excel — not the
    numbers the quote was first generated with.

What it reads from the Excel:
    - "Meta" tab  -> Project, Quote Ref, Customer (used for the order header).
    - "Quote" tab -> every line's Area, Room, System, Item, SKU, Qty, Unit Price,
      and the three labour hour columns (1st Fix / 2nd Fix / Prog), exactly as they
      stand after editing. System is checked against Odoo's live System list.

What it does:
    - Re-resolves each equipment line to its live Odoo product id (by SKU, then by
      name) — only the id is taken from Odoo; the PRICE and QTY come from the Excel,
      so any manual price override you made is preserved.
    - Re-aggregates labour from the (possibly edited) hour columns into the three
      Odoo labour service lines.
    - Writes a clean Odoo import CSV (same exact format build_quote.py uses).
    - Anything it can't resolve to an Odoo product id is left out of the import and
      reported, so nothing imports against the wrong product.

Needs the same Odoo env vars as build_quote.py (ODOO_URL/DB/LOGIN/KEY). If Odoo
isn't reachable it can't attach product ids, so it stops rather than writing an
import file that wouldn't load.
"""
import sys
import os
import json

from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import build_quote as bq
try:
    import odoo_client
except Exception:
    odoo_client = None


def _norm(v):
    return "" if v is None else str(v).strip()


def _num(v):
    if v is None or v == "":
        return 0.0
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def read_meta(wb):
    meta = {"project": "Untitled Project", "quote_ref": "", "customer": ""}
    if "Meta" not in wb.sheetnames:
        return meta
    ws = wb["Meta"]
    label_to_key = {v.lower(): k for k, v in bq.META_LABELS.items()}
    for row in ws.iter_rows(min_row=1, max_col=2, values_only=True):
        label = _norm(row[0]).lower()
        if label in label_to_key:
            meta[label_to_key[label]] = _norm(row[1])
    if not meta["project"]:
        meta["project"] = "Untitled Project"
    return meta


# Row markers in column A that mean "the line-item table has ended"
STOP_MARKERS = ("TOTALS", "GRAND TOTAL", "EQUIPMENT MARGIN", "PRICING:")


def read_quote_lines(wb):
    """Read the edited 'Quote' tab into line dicts build_odoo_import understands."""
    if "Quote" not in wb.sheetnames:
        raise SystemExit("ERROR: no 'Quote' tab found in the workbook.")
    ws = wb["Quote"]

    # find the header row (the row containing 'Item' and 'Qty')
    header_row = None
    header = {}
    for r in range(1, 6):
        vals = {c: _norm(ws.cell(row=r, column=c).value) for c in range(1, ws.max_column + 1)}
        names = {v: c for c, v in vals.items() if v}
        if "Item" in names and "Qty" in names:
            header_row = r
            header = names
            break
    if header_row is None:
        raise SystemExit("ERROR: couldn't find the header row (need 'Item' and 'Qty' columns).")

    def col(name):
        return header.get(name)

    need = ["Item", "SKU", "Qty", "Unit Price", "1st Fix Hrs", "2nd Fix Hrs", "Prog Hrs"]
    missing = [n for n in need if col(n) is None]
    if missing:
        raise SystemExit(f"ERROR: the Quote tab is missing expected columns: {missing}")

    lines = []
    for r in range(header_row + 1, ws.max_row + 1):
        # the totals/footer markers are written in the leftmost column
        first_cell = _norm(ws.cell(row=r, column=1).value)
        if first_cell.upper().startswith(STOP_MARKERS):
            break
        cat = _norm(ws.cell(row=r, column=col("Category")).value) if col("Category") else ""
        # Area / Room / System each have their own column. Older workbooks only
        # had a combined "Room / Area" column — read that as Room.
        area = _norm(ws.cell(row=r, column=col("Area")).value) if col("Area") else ""
        room_col = col("Room") or col("Room / Area")
        room = _norm(ws.cell(row=r, column=room_col).value) if room_col else ""
        system = _norm(ws.cell(row=r, column=col("System")).value) if col("System") else ""
        item = _norm(ws.cell(row=r, column=col("Item")).value)
        sku = _norm(ws.cell(row=r, column=col("SKU")).value)
        qty = _num(ws.cell(row=r, column=col("Qty")).value)
        if not item and not qty:
            continue  # blank spacer row
        if sku in ("", "-"):
            sku = ""
        lines.append({
            "query": sku or item,
            "sku": sku,
            "item": item,
            "category": cat,
            "area": area,
            "room": room,
            "system": system,
            "qty": qty,
            "unit_price": _num(ws.cell(row=r, column=col("Unit Price")).value),
            "unit_cost": _num(ws.cell(row=r, column=col("Unit Cost")).value) if col("Unit Cost") else 0.0,
            "fix1_h": _num(ws.cell(row=r, column=col("1st Fix Hrs")).value),
            "fix2_h": _num(ws.cell(row=r, column=col("2nd Fix Hrs")).value),
            "prog_h": _num(ws.cell(row=r, column=col("Prog Hrs")).value),
        })
    return lines


def read_fees(wb):
    """Read the FEE: rows (project management / design hours in the Qty column)."""
    out = {"project_management": 0.0, "design": 0.0}
    ws = wb["Quote"]
    qty_col = None
    for r in range(1, 6):
        for c in range(1, ws.max_column + 1):
            if _norm(ws.cell(row=r, column=c).value) == "Qty":
                qty_col = c
    if qty_col is None:
        return out
    for r in range(1, ws.max_row + 1):
        label = _norm(ws.cell(row=r, column=1).value).upper()
        if label.startswith("FEE: PROJECT MANAGEMENT"):
            out["project_management"] = _num(ws.cell(row=r, column=qty_col).value)
        elif label.startswith("FEE: DESIGN"):
            out["design"] = _num(ws.cell(row=r, column=qty_col).value)
    return out


def resolve_ids(raw_lines):
    """Attach a live Odoo product id to each line (by SKU, then name)."""
    if odoo_client is None or not odoo_client.available():
        raise SystemExit("ERROR: Odoo not reachable (no ODOO_KEY). Can't build an import "
                         "file without live product ids — set the env vars and retry.")
    try:
        od = odoo_client.Odoo()
    except Exception as e:
        raise SystemExit(f"ERROR: Odoo connection failed ({e}).")

    try:
        sys_options = odoo_client.system_options(od)
    except Exception:
        sys_options = []

    out = []
    for L in raw_lines:
        odoo_id = None
        matched_name = L["item"] or L["query"]
        q = L["sku"] or L["item"]
        try:
            prod, _mt = odoo_client.lookup_product(od, q)
        except Exception:
            prod = None
        if prod:
            odoo_id = prod["id"]
            matched_name = L["item"] or prod.get("name") or q
        out.append({
            "query": L["query"], "matched_name": matched_name, "sku": L["sku"],
            "item": L["item"], "category": L.get("category", ""),
            "area": L.get("area", ""), "room": L.get("room", ""),
            "system_key": (bq.normalise_system(L.get("system"), sys_options) or ("", ""))[0],
            "system_in": L.get("system", ""),
            "qty": L["qty"], "unit_price": L["unit_price"], "unit_cost": L.get("unit_cost", 0.0),
            "fix1_h": L["fix1_h"], "fix2_h": L["fix2_h"], "prog_h": L["prog_h"],
            "odoo_id": odoo_id,
            "price_source": "from edited Excel" if odoo_id else "not found in Odoo",
        })
    return out


def write_new_products_file(lines, out_path):
    """For lines with no Odoo match, seed a new_products.json from the Excel so a
    person can review/complete it and run create_products.py. Returns the list of
    unmatched (new) items, or [] if none."""
    new_items = []
    for L in lines:
        if L["odoo_id"]:
            continue
        if not (L.get("item") or L.get("sku")) or not L["qty"]:
            continue
        new_items.append({
            "name": L.get("item") or L["query"],
            "sku": L["sku"],
            "list_price": L["unit_price"],
            "standard_price": L.get("unit_cost", 0.0),
            "category": L.get("category", ""),
            "type": "consu",
        })
    if new_items:
        with open(out_path, "w") as f:
            json.dump({"products": new_items}, f, indent=2)
    return new_items


def main():
    if len(sys.argv) not in (2, 3):
        print("Usage: python refresh_import.py <edited_quote.xlsx> [odoo_import.csv]")
        sys.exit(1)
    xlsx_path = sys.argv[1]
    if len(sys.argv) == 3:
        out_path = sys.argv[2]
    else:
        stem = xlsx_path[:-5] if xlsx_path.lower().endswith(".xlsx") else xlsx_path
        out_path = stem + "_OdooImport.csv"

    wb = load_workbook(xlsx_path, data_only=True)
    meta = read_meta(wb)
    raw_lines = read_quote_lines(wb)
    # Labour products (CLH / INSTL / PROGL) are generated from hours only. If one
    # was typed in as an equipment line (or a banned LAB-* code), drop it so
    # labour is never double-counted or booked against another SKU.
    kept = []
    for L in raw_lines:
        if bq.is_labour_item(L.get("sku"), L.get("item"), L.get("query")) or bq.labour_category(L.get("category")):
            print(f"WARNING: row '{L.get('item') or L.get('query')}' is a labour line - skipped. "
                  f"Labour is billed only via CLH / INSTL / PROGL from the hours columns.")
            continue
        kept.append(L)
    raw_lines = kept
    lines = resolve_ids(raw_lines)

    _, _, _, _, _, rates = bq.load_reference()
    labour_products, labour_note = bq.resolve_labour_products(rates)
    fees = read_fees(wb)
    fee_products, fee_note = bq.resolve_fee_products()

    rows, skipped = bq.build_odoo_import(
        lines, meta["project"], meta["quote_ref"], meta["customer"], rates, labour_products,
        fees, fee_products
    )
    bq.write_odoo_import(out_path, rows)

    # any line with no Odoo match is a candidate NEW product to add to Odoo
    new_path = (xlsx_path[:-5] if xlsx_path.lower().endswith(".xlsx") else xlsx_path) + "_NewProducts.json"
    new_items = write_new_products_file(lines, new_path)

    n_equip = sum(1 for l in lines if l["odoo_id"])
    bad_sys = [l for l in lines if l["odoo_id"] and l["system_in"] and not l["system_key"]]
    no_sys = [l for l in lines if l["odoo_id"] and not l["system_in"]]
    print(f"Refreshed {out_path} from {os.path.basename(xlsx_path)}")
    print(f"  Project: {meta['project']} | Ref: {meta['quote_ref'] or '(none)'} | "
          f"Customer: {meta['customer'] or '(set at import)'}")
    print(f"  {len(lines)} lines read, {n_equip} equipment lines matched to Odoo, "
          f"{len(rows)} importable rows written")
    print(f"  labour: {labour_note}")
    print(f"  fees: PM {fees['project_management']} hrs, design {fees['design']} hrs — {fee_note}")
    if not meta["customer"]:
        print("  NOTE: no Customer on the Meta tab — set the Customer column in Odoo at import.")
    for name, reason in skipped:
        print(f"  SKIPPED from import: {name} — {reason}")
    for l in bad_sys:
        print(f'  SYSTEM NOT VALID (imported blank): {l["matched_name"]} — "{l["system_in"]}" '
              f"is not an Odoo System type")
    if no_sys:
        print(f"  {len(no_sys)} equipment line(s) have no System set — they import with System blank:")
        for l in no_sys:
            print(f"    - {l['matched_name']}")
    if new_items:
        print(f"\n  {len(new_items)} line(s) are NOT in Odoo yet and were left out of the import.")
        print(f"  To add them to the Odoo pricelist so they import, review/complete")
        print(f"  {os.path.basename(new_path)} (name, SKU, sell price, cost, category, type),")
        print(f"  run:  python create_products.py {os.path.basename(new_path)}")
        print(f"  then re-run this refresh. New (not-in-Odoo) items:")
        for ni in new_items:
            print(f"    - {ni['name']} [{ni['sku'] or 'no SKU'}] sell={ni['list_price']} cost={ni['standard_price']}")


if __name__ == "__main__":
    main()
