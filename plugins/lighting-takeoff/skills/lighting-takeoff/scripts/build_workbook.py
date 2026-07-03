#!/usr/bin/env python3
"""
build_workbook.py — turn a room-level lighting take-off JSON into a D-One xlsx.

Each floor sheet is a room-by-room matrix: one row per room, with columns for
every light-fitting type, LED strip runs, switch points, every circuit type, and
suggested sensors, finished with a bold TOTAL row. A Summary sheet then totals
each item across floors. Recording counts per room (not just per floor) makes the
take-off auditable — you can see exactly which room each fitting was counted in.

Usage:
    python build_workbook.py counts.json OUTPUT.xlsx

counts.json schema (room-level, preferred):
{
  "project": "12 Mountain Road",
  "floors": [
    {
      "name": "First Floor",
      "rooms": [
        {
          "name": "Master Bedroom",
          "lights": { "Recessed downlight": 6, "Wall mounted light point": 4 },
          "led_strips": 2,
          "switch_points": 2,
          "circuits": { "Downlight circuit": 1, "Wall-light circuit": 1, "LED-strip circuit": 1 },
          "sensors": 0
        },
        {
          "name": "Master en-suite",
          "lights": { "Recessed downlight": 4 },
          "led_strips": 1, "switch_points": 1,
          "circuits": { "Downlight circuit": 1, "LED-strip circuit": 1 },
          "sensors": 1
        }
      ],
      "notes": "optional free text"
    }
  ]
}

Conventions:
- `lights` keys are free-form fitting names from the plan legend.
- `switch_points` = number of wall plate locations (each becomes a keypad), one
  per location regardless of how many switches/gangs sit there.
- `led_strips` = number of distinct LED runs.
- `circuits` = {circuit type: count}; a circuit is connected fittings/runs
  switched together (so LED runs >= LED-strip circuits).
- `sensors` = suggested motion sensors for that room (sculleries/ensuites/garages).
Floor and project totals are computed from the rooms.

Backward-compatible: a floor may instead carry floor-level `lights`/`led_strips`/
`switch_points`/`circuits`/`sensor_rooms` (no `rooms`); it is then rendered as a
single pseudo-room called "(floor total)".
"""
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BLUE = "1379C9"
DARK = "0B3C5D"
LIGHT = "E8F1FB"
GREY = "ECECEC"

H_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(name="Calibri", bold=True, color=DARK, size=16)
SUB_FONT = Font(name="Calibri", italic=True, color="666666", size=9)
BOLD = Font(name="Calibri", bold=True, size=11)
NORM = Font(name="Calibri", size=11)
H_FILL = PatternFill("solid", fgColor=BLUE)
ALT_FILL = PatternFill("solid", fgColor=LIGHT)
TOT_FILL = PatternFill("solid", fgColor=GREY)
THIN = Side(style="thin", color="D0D0D0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CTR = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")
ROT = Alignment(horizontal="center", vertical="bottom", text_rotation=90)


def norm_rooms(floor):
    """Return a list of room dicts whether the floor is room-level or floor-level."""
    if floor.get("rooms"):
        return floor["rooms"]
    # legacy floor-level -> one pseudo-room
    circ = floor.get("circuits", {})
    if not isinstance(circ, dict):
        circ = {"Lighting circuit": circ or 0}
    sensors = sum(s.get("sensors", 0) or 0 for s in floor.get("sensor_rooms", []) or [])
    return [{"name": "(floor total)", "lights": floor.get("lights", {}) or {},
             "led_strips": floor.get("led_strips", 0),
             "keypads": floor.get("keypads",
                                  floor.get("switch_points", floor.get("switch_positions", 0))),
             "circuits": circ, "sensors": sensors}]


def collect_types(floors):
    light_types, circ_types = [], []
    for fl in floors:
        for r in norm_rooms(fl):
            for k in (r.get("lights") or {}):
                if k not in light_types:
                    light_types.append(k)
            c = r.get("circuits") or {}
            if isinstance(c, dict):
                for k in c:
                    if k not in circ_types:
                        circ_types.append(k)
    return light_types, circ_types


def room_vals(room, light_types, circ_types):
    lights = room.get("lights") or {}
    circ = room.get("circuits") or {}
    if not isinstance(circ, dict):
        circ = {}
    lv = [lights.get(t, 0) for t in light_types]
    fittings = sum(v or 0 for v in lv)
    led = room.get("led_strips", 0) or 0
    # one control location per spot, whether drawn as a switch bank or a keypad
    kp = room.get("keypads", room.get("switch_points", room.get("switch_positions", 0))) or 0
    cv = [circ.get(t, 0) for t in circ_types]
    ctot = sum(v or 0 for v in cv)
    sens = room.get("sensors", 0) or 0
    return lv, fittings, led, kp, cv, ctot, sens


def build_floor_sheet(wb, floor, light_types, circ_types):
    ws = wb.create_sheet(title=floor["name"][:31])
    cols = (["Room"] + light_types + ["Fittings", "LED runs", "Keypads"]
            + circ_types + ["Circuits", "Sensors"])
    ncol = len(cols)

    ws.cell(row=1, column=1, value=floor["name"]).font = TITLE_FONT
    ws.cell(row=2, column=1,
            value="Room-by-room lighting take-off — visual counts, verify before ordering").font = SUB_FONT

    hrow = 4
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=hrow, column=i, value=c)
        cell.fill = H_FILL; cell.font = H_FONT; cell.border = BORDER
        cell.alignment = LEFT if i == 1 else ROT
    ws.row_dimensions[hrow].height = 96

    totals = [0] * (ncol - 1)
    r = hrow + 1
    for ridx, room in enumerate(norm_rooms(floor)):
        lv, fittings, led, kp, cv, ctot, sens = room_vals(room, light_types, circ_types)
        rowvals = lv + [fittings, led, kp] + cv + [ctot, sens]
        a = ws.cell(row=r, column=1, value=room.get("name", "?"))
        a.font = NORM; a.alignment = LEFT; a.border = BORDER
        if ridx % 2 == 1:
            a.fill = ALT_FILL
        for j, v in enumerate(rowvals, 2):
            cell = ws.cell(row=r, column=j, value=v)
            cell.font = NORM; cell.alignment = CTR; cell.border = BORDER
            if ridx % 2 == 1:
                cell.fill = ALT_FILL
            totals[j - 2] += v or 0
        r += 1

    t = ws.cell(row=r, column=1, value="TOTAL")
    t.font = BOLD; t.alignment = LEFT; t.border = BORDER; t.fill = TOT_FILL
    for j, v in enumerate(totals, 2):
        cell = ws.cell(row=r, column=j, value=v)
        cell.font = BOLD; cell.alignment = CTR; cell.border = BORDER; cell.fill = TOT_FILL

    ws.column_dimensions["A"].width = 26
    for i in range(2, ncol + 1):
        ws.column_dimensions[get_column_letter(i)].width = 8
    ws.freeze_panes = "B5"

    if floor.get("notes"):
        ws.cell(row=r + 2, column=1, value="Notes: " + floor["notes"]).font = SUB_FONT

    # floor aggregate for the summary
    agg = {"name": floor["name"], "light_types": {}, "fittings": 0, "led": 0,
           "keypads": 0, "circ_types": {}, "circuits": 0, "sensors": 0}
    for room in norm_rooms(floor):
        lv, fittings, led, kp, cv, ctot, sens = room_vals(room, light_types, circ_types)
        for t_, v in zip(light_types, lv):
            agg["light_types"][t_] = agg["light_types"].get(t_, 0) + (v or 0)
        for t_, v in zip(circ_types, cv):
            agg["circ_types"][t_] = agg["circ_types"].get(t_, 0) + (v or 0)
        agg["fittings"] += fittings; agg["led"] += led; agg["keypads"] += kp
        agg["circuits"] += ctot; agg["sensors"] += sens
    return agg


def build_summary(wb, project, aggs, light_types, circ_types):
    ws = wb.create_sheet(title="Summary")
    wb.move_sheet("Summary", -(len(wb.sheetnames) - 1))
    floors = [a["name"] for a in aggs]

    ws.column_dimensions["A"].width = 42
    for i in range(len(floors) + 1):
        ws.column_dimensions[get_column_letter(2 + i)].width = 14

    ws.cell(row=1, column=1, value=(project or "Lighting Take-off") + " — Summary").font = TITLE_FONT
    ws.cell(row=2, column=1,
            value="Visual estimates from the plans — sanity-check before ordering. See each floor tab for the room-by-room breakdown.").font = SUB_FONT
    row = 4
    for i, name in enumerate(["Item"] + floors + ["Total"], 1):
        cell = ws.cell(row=row, column=i, value=name)
        cell.fill = H_FILL; cell.font = H_FONT; cell.border = BORDER
        cell.alignment = LEFT if i == 1 else CTR
    row += 1

    def srow(label, getter, bold=False):
        nonlocal row
        a = ws.cell(row=row, column=1, value=label)
        a.alignment = LEFT; a.font = BOLD if bold else NORM; a.border = BORDER
        if bold:
            a.fill = TOT_FILL
        tot = 0
        for j, ag in enumerate(aggs):
            v = getter(ag); tot += v or 0
            c = ws.cell(row=row, column=2 + j, value=v)
            c.alignment = CTR; c.font = BOLD if bold else NORM; c.border = BORDER
            if bold:
                c.fill = TOT_FILL
        c = ws.cell(row=row, column=2 + len(aggs), value=tot)
        c.alignment = CTR; c.font = BOLD; c.border = BORDER; c.fill = TOT_FILL
        row += 1

    for t in light_types:
        srow(t, lambda a, t=t: a["light_types"].get(t, 0))
    srow("Total light fittings", lambda a: a["fittings"], bold=True)
    srow("LED strip runs", lambda a: a["led"])
    srow("Keypads (switch-bank / keypad locations)", lambda a: a["keypads"])
    for ct in circ_types:
        srow(ct, lambda a, ct=ct: a["circ_types"].get(ct, 0))
    srow("Total lighting circuits", lambda a: a["circuits"], bold=True)
    srow("Suggested motion sensors", lambda a: a["sensors"], bold=True)


def main():
    data = json.load(open(sys.argv[1]))
    out = sys.argv[2]
    light_types, circ_types = collect_types(data["floors"])
    wb = Workbook(); wb.remove(wb.active)
    aggs = [build_floor_sheet(wb, fl, light_types, circ_types) for fl in data["floors"]]
    build_summary(wb, data.get("project"), aggs, light_types, circ_types)
    wb.save(out)
    print("wrote", out)


if __name__ == "__main__":
    main()
